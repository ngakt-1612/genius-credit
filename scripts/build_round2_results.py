#!/usr/bin/env python3
"""Gộp kết quả Daily Quota vòng 2 → results-round2.json + thư mục ảnh bằng chứng + xlsx kết quả.

- TC giữ kết quả 22/10 (r2=CARRY): status đã chấm lại theo Expected mới + ảnh 22/10.
- TC chạy vòng 2: dòng `VERDICT|N<số>[-phần]|STATUS|...` trong runs/*.verdicts (lượt sau thắng),
  ảnh là attachment có tên chứa `N<số>` trong runs/*-shots/manifest.json.
- TC Gate chưa chạy: «Chờ» + lý do.
File TC gốc chỉ đọc; xuất `TC_IIP055.1.16_Daily_Quota_Topup_ket_qua_v2.xlsx`.
"""
import glob
import json
import os
import re
import shutil
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

HERE = Path(__file__).resolve().parent
TC_FILE = Path("~/Downloads/TC_IIP055.1.16_Daily_Quota_Topup.xlsx").expanduser()
OUT_XLSX = Path("~/Downloads/TC_IIP055.1.16_Daily_Quota_Topup_ket_qua_v2.xlsx").expanduser()
EVID = Path("~/Downloads/Genius-2.9.0-Daily-Quota-round2-bang-chung").expanduser()
STATUS = {"PASS": "PASS", "FAIL": "FAIL", "NA": "N/A", "N/A": "N/A", "CHECK": "Cần xem", "MISMATCH": "Lệch spec", "WAIT": "Chờ"}
FILL = {"PASS": "C6EFCE", "FAIL": "FFC7CE", "Lệch spec": "FFEB9C", "Cần xem": "DDEBF7", "Chờ": "EDEDED", "Pending": "FCE4D6", "N/A": "EDEDED"}
GATE_REASON = {
    "X-BE": "Chờ BE/Infra: cần giả lập phía server (vendor lỗi/timeout, webhook, gọi API trực tiếp, proxy) — theo luật không tự giả lỗi server",
    "S9-STOREKIT": "Chờ tài nguyên: sandbox trên máy tự xác nhận mua → cần tài khoản sandbox khác (huỷ sheet/thẻ lỗi/Ask to Buy) hoặc StoreKit Configuration",
    "S9-2DEV": "Chờ tài nguyên: cần máy thứ 2 cùng Apple ID",
    "S7-CMS": "Chờ PO: đổi số lượt/ngày (Q) trên CMS",
    "S8-RESET": "Pending (user 24/10 17:41): cần khung 06:50–07:10 sáng (mốc reset 00:00 UTC) — chưa test",
}
# Lý do Chờ cụ thể cho TC chưa viết test (soát 25/10)
WAIT_REASON = {
    "25": "Chờ tài nguyên: cần máy A + máy B cùng Apple ID", "37": "Chờ tài nguyên: cần máy A + máy B cùng Apple ID",
    "52": "Chờ tài nguyên: cần máy A + máy B cùng Apple ID", "133": "Chờ tài nguyên: cần máy B chưa đăng nhập",
    "45": "Chờ tài nguyên: cần ảnh mẫu bị chặn NSFW do team cung cấp (không tự tạo nội dung NSFW)",
    "57": "Chưa chạy: đổi ngày giờ hệ thống trên máy test (ảnh hưởng sandbox & lượt chạy khác) — cần người đồng ý",
    "83": "Chưa chạy: cần sub grandfather gia hạn qua mốc (mốc RC nằm giữa 2 lần gia hạn sandbox) — dựng được bằng RC ép trong code, chưa làm",
    "122": "Chưa chạy: như N083 (chu kỳ đang trả bắt đầu sau mốc, sub đầu trước mốc)",
    "128": "Chờ người: đăng nhập Apple trong app cần mật khẩu Apple ID",
    "134": "Không tự làm: Erase All Content and Settings xoá máy test — cần người quyết",
    "136": "Phần app: key daily_quota_amount đã gỡ khỏi RemoteConfigKeys.swift:216–218 (đọc code); 413=15 cần PO đặt Q=15 trên CMS — Chờ PO",
    "158": "Chờ BE/Infra: cần proxy bắt request offering", "159": "Chờ BE/Infra: cần proxy bắt request lúc launch",
    "171": "Chưa chạy: Restore Purchases trên app cài mới (chạy được — chưa làm)",
    "172": "Chờ server: gen video Seedance (I2V) vào từ tab Home — Home lỗi server 25/10",
}
# Case phải chờ lâu, user cho Pending 24/10 17:41
PENDING_LONG_WAIT = {
    "48": "Pending: kill app giữa gen rồi chờ BE auto-release 10′ — user chỉ chạy case không chờ lâu (24/10 17:41)",
    "49": "Pending: như N048 (đo lúc 9′)", "50": "Pending: như N048", "100": "Pending: toast 483 sau auto-release 10′ (như N048)",
}


def round2_verdicts():
    """{N°: [(phần, status, actual, run)]} — lượt sau thắng cho cùng id."""
    latest = {}
    for f in sorted(glob.glob(str(HERE / "runs/*.verdicts")), key=os.path.getmtime):
        run = Path(f).stem
        for line in Path(f).read_text(encoding="utf-8").splitlines():
            m = re.match(r"VERDICT\|N(\d+)(-[\w-]+)?\|(\w+/?\w*)\|(.*)", line)
            if m:
                latest[(m.group(1).lstrip("0"), m.group(2) or "")] = (STATUS.get(m.group(3), m.group(3)), m.group(4), run)
    out = defaultdict(list)
    for (tc, part), v in latest.items():
        out[tc].append((part,) + v)
    return out


def merge_parts(parts):
    sts = [p[1] for p in parts if not p[0].endswith("pre")] or [p[1] for p in parts]
    for s in ("FAIL", "Lệch spec", "Cần xem", "N/A"):
        if s in sts:
            return s
    return "PASS" if all(s == "PASS" for s in sts) else sts[-1]


def round2_shots():
    shots = defaultdict(list)
    for m in glob.glob(str(HERE / "runs/*-shots/manifest.json")):
        for t in json.load(open(m)):
            for a in t.get("attachments", []):
                for tc in re.findall(r"\bN(\d{3})\b", a["suggestedHumanReadableName"]):
                    shots[tc.lstrip("0")].append(os.path.join(os.path.dirname(m), a["exportedFileName"]))
    return shots


def main():
    matrix = json.load(open(HERE / "quota-matrix.json"))
    r2v, r2s = round2_verdicts(), round2_shots()
    manual = json.load(open(HERE / "manual-regrade.json")) if (HERE / "manual-regrade.json").exists() else {}
    if EVID.exists():
        shutil.rmtree(EVID)
    results = {}
    for c in matrix:
        n = c["id"]
        if c["r2"] == "CARRY":
            st = c.get("new_status") or c["old_status"]
            actual = c["old_actual"]
            note = f"KQ autotest 22/10 (TC cũ #{c['old_id']}, lượt {c['old_run'] or 'đọc code'})"
            if c.get("regrade_note"):
                note += f" · chấm lại theo Expected mới: {c['regrade_note']}"
            files = [e["path"] for e in c.get("evidence", [])]
        elif n in r2v:
            parts = r2v[n]
            st = merge_parts(parts)
            if n in manual:
                st = manual[n]["status"]
            actual = " ‖ ".join(f"{p[0].lstrip('-') or 'KQ'}: {p[2]}" for p in sorted(parts))
            note = "Autotest vòng 2 24/10 · lượt " + ", ".join(sorted({p[3] for p in parts}))
            if n in manual:
                note += " · CHẤM TAY: " + manual[n]["why"]
            files = r2s.get(n, [])
        elif n in manual:
            st, actual, note, files = manual[n]["status"], "", "Chấm từ bằng chứng/code: " + manual[n]["why"], []
        else:
            st, actual = "Chờ", ""
            note = PENDING_LONG_WAIT.get(n) or WAIT_REASON.get(n) or GATE_REASON.get(c["suite"], "Chưa chạy — xem handoff")
            if n in PENDING_LONG_WAIT or c["suite"] == "S8-RESET":
                st = "Pending"
            files = []
        dest = EVID / f"N{int(n):03d}"
        kept = []
        for i, f in enumerate(files[:6], 1):
            if os.path.exists(f):
                dest.mkdir(parents=True, exist_ok=True)
                d = dest / f"N{int(n):03d}-{i}.png"
                shutil.copy2(f, d)
                kept.append(str(d))
        results[n] = {"status": st, "actual": actual, "note": note, "evidence": kept, "suite": c["suite"]}
    json.dump(results, open(HERE / "results-round2.json", "w"), ensure_ascii=False, indent=1)

    wb = openpyxl.load_workbook(TC_FILE)
    ws = wb["TC_Quota_Topup"]
    head = {"I": "PASS/FAIL", "J": "Kết quả thực tế", "K": "Ghi chú / nguồn kết quả", "L": "Ảnh bằng chứng"}
    for col, title in head.items():
        ws[f"{col}2"] = title
        ws[f"{col}2"].font = Font(bold=True)
    for col, w in (("I", 12), ("J", 70), ("K", 45), ("L", 22)):
        ws.column_dimensions[col].width = w
    for c in matrix:
        r, res = c["row"], results[c["id"]]
        ws[f"I{r}"] = res["status"]
        ws[f"I{r}"].fill = PatternFill("solid", fgColor=FILL.get(res["status"], "FFFFFF"))
        ws[f"J{r}"] = res["actual"]
        ws[f"K{r}"] = res["note"]
        for col in "JK":
            ws[f"{col}{r}"].alignment = Alignment(wrap_text=True, vertical="top")
        if res["evidence"]:
            ws[f"L{r}"] = f"{len(res['evidence'])} ảnh — mở thư mục"
            ws[f"L{r}"].hyperlink = Path(res["evidence"][0]).parent.as_uri()
            ws[f"L{r}"].font = Font(color="0563C1", underline="single")
    wb.save(OUT_XLSX)
    tally = defaultdict(int)
    for v in results.values():
        tally[v["status"]] += 1
    print(json.dumps({"xlsx": str(OUT_XLSX), "evidence": str(EVID), **tally}, ensure_ascii=False))


if __name__ == "__main__":
    sys.exit(main())
